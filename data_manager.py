from datetime import datetime
from pathlib import Path
from time import sleep
import secrets

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SHEET_DATA = "注册信息"
SHEET_HELP = "填写说明"
SHEET_POLICY = "政策模板"

POLICY_COL = {
    "Shipping policy": "运费政策",
    "Legal notice": "法律声明",
    "Terms of sale": "销售条款",
}

SAMPLE_POLICIES = [
    (
        "Shipping policy",
        "Shipping Policy for {shop_name}\n\n"
        "Orders are processed after payment is confirmed. Delivery times vary by destination and carrier. "
        "You will receive tracking information by email when your order ships.\n\n"
        "Questions about shipping: {email}",
    ),
    (
        "Shipping policy",
        "{shop_name} Shipping\n\n"
        "We dispatch in-stock items as soon as possible after checkout. Shipping costs and estimated dates are shown at checkout. "
        "If a package is delayed, contact the carrier with your tracking number first.\n\n"
        "Need help? Email {email}",
    ),
    (
        "Shipping policy",
        "Delivery information — {shop_name}\n\n"
        "Standard shipping is available to the destinations we support at checkout. "
        "Please allow extra time for remote areas and peak seasons. Lost or damaged parcels should be reported promptly.\n\n"
        "Contact: {email}",
    ),
    (
        "Legal notice",
        "Legal Notice for {shop_name}\n\n"
        "{shop_name} operates this online store. For legal or privacy questions, contact {email}. "
        "Please also review our return, shipping, and terms of service policies.",
    ),
    (
        "Legal notice",
        "Notice — {shop_name}\n\n"
        "This website is operated by {shop_name}. Content is provided for store customers. "
        "If you believe any information is incorrect, write to {email}.",
    ),
    (
        "Legal notice",
        "{shop_name} legal information\n\n"
        "Use of this store is subject to our published policies. "
        "For company or legal correspondence, email {email}.",
    ),
    (
        "Terms of sale",
        "Terms of Sale — {shop_name}\n\n"
        "By placing an order you agree to these terms. Product descriptions and prices are shown at checkout. "
        "Payment is due when you order. For order issues contact {email}.",
    ),
    (
        "Terms of sale",
        "{shop_name} Sale Terms\n\n"
        "All sales are for the items listed on the product page at the price displayed at checkout. "
        "We may cancel an order if an item cannot be fulfilled and will notify you by email.\n\n"
        "Questions: {email}",
    ),
    (
        "Terms of sale",
        "Purchase terms for {shop_name}\n\n"
        "An order is an offer to buy. We accept that offer when we confirm or ship the order. "
        "Returns follow our return policy. Contact {email} for billing questions.",
    ),
]

HEADERS = [
    "AdsPower环境ID",
    "店铺名",
    "邮箱",
    "邮箱密码",
    "IMAP服务器",
    "Shopify密码",
    "联系人姓名",
    "国家",
    "地址",
    "城市",
    "州/省",
    "邮编",
    "电话",
    "产品CSV",
    "横幅图片",
    "运费政策",
    "法律声明",
    "销售条款",
    "状态",
    "店铺后台",
    "装修状态",
    "失败原因",
    "完成时间",
]

REQUIRED_FIELDS = ["AdsPower环境ID", "店铺名", "邮箱", "联系人姓名"]

HELP_ROWS = [
    ["列名", "必填", "说明"],
    ["AdsPower环境ID", "是", "AdsPower 窗口列表里的 ID（不是序号）。点开环境详情可复制，例如 k1x2abc。一店一环境。"],
    ["店铺名", "是", "店铺名称，将生成 店铺名.myshopify.com，只能用字母数字和连字符。"],
    ["邮箱", "是", "你自己的、能收信的真实邮箱，用于 Shopify 账号。"],
    ["邮箱密码", "否", "该邮箱的 IMAP 密码或应用专用密码。填写后脚本尝试自动读验证码；留空则注册时在终端暂停，由你手动填。"],
    ["IMAP服务器", "否", "覆盖 config.ini 的默认值。Outlook: outlook.office365.com；Gmail: imap.gmail.com。"],
    ["Shopify密码", "否", "你为该店设置的登录密码。留空则脚本生成一串强密码并写回本列。"],
    ["联系人姓名", "是", "店主真实姓名，需与你向 Shopify 申报的身份一致。"],
    ["国家", "否", "经营所在国家，建议填英文，例如 United States / China。"],
    ["地址", "否", "街道地址。"],
    ["城市", "否", "城市。"],
    ["州/省", "否", "州或省，美国填两位州代码更稳，例如 CA。"],
    ["邮编", "否", "邮政编码。"],
    ["电话", "否", "联系电话，含区号，例如 +1 4155550123。"],
    ["产品CSV", "否", "注册完成后要导入的 Shopify 产品 CSV 完整路径。"],
    ["横幅图片", "否", "主题首页横幅图片完整路径。留空则尝试用 Shopify 主题 AI 生成。"],
    ["运费政策", "否", "本店 Shipping policy 正文。留空则从「政策模板」表随机抽一条。可用 {shop_name} {email}。"],
    ["法律声明", "否", "本店 Legal notice 正文。留空则从「政策模板」表随机抽一条。"],
    ["销售条款", "否", "本店 Terms of sale 正文。留空则从「政策模板」表随机抽一条。"],
    ["状态", "系统写回", "待处理 / 成功 / 失败。退货规则、书面政策都做完才算成功；主题不用做。只注册成功不算。"],
    ["店铺后台", "系统写回", "注册成功后的后台地址。"],
    ["装修状态", "系统写回", "导入产品、政策等后续步骤的结果。"],
    ["失败原因", "系统写回", "失败时的原因。"],
    ["完成时间", "系统写回", "处理完成时间。"],
    [],
    ["使用前请确认"],
    ["1. AdsPower 客户端已打开，本地 API 为开启状态（默认 http://127.0.0.1:50325）。"],
    ["2. 每个店铺对应一个已经创建好的 AdsPower 环境，把 ID 填进表格。"],
    ["3. 只填写你本人/公司真实、有权使用的邮箱和身份信息。"],
    ["4. 填好本表后运行：python main.py"],
]


def ensure_extra_columns(excel_path: str):
    extra = ["产品CSV", "横幅图片", "装修状态", "运费政策", "法律声明", "销售条款"]
    path = Path(excel_path)
    if not path.exists():
        return
    wb = load_workbook(excel_path)
    if SHEET_DATA not in wb.sheetnames:
        return
    ws = wb[SHEET_DATA]
    headers = [_cell_str(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    changed = False
    for name in extra:
        if name not in headers:
            col = len(headers) + 1
            ws.cell(1, col, name)
            headers.append(name)
            changed = True
    if changed:
        _save_workbook(wb, excel_path)
        print("已在 Excel 中补上政策/装修相关列。")
    ensure_policy_sheet(excel_path)


def ensure_template(excel_path: str):
    """若注册表不存在，则生成带表头和说明的模板。"""
    if Path(excel_path).exists():
        ensure_extra_columns(excel_path)
        return
    create_template(excel_path)
    print(f"已生成数据模板：{excel_path}，请先按「填写说明」填入真实资料后再运行。")


def create_template(excel_path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_DATA
    help_ws = wb.create_sheet(SHEET_HELP)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    required_fill = PatternFill("solid", fgColor="1F4E79")
    optional_fill = PatternFill("solid", fgColor="375623")
    system_fill = PatternFill("solid", fgColor="595959")
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    required = set(REQUIRED_FIELDS)
    system = {"状态", "店铺后台", "失败原因", "完成时间"}

    for col, name in enumerate(HEADERS, 1):
        cell = ws.cell(1, col, name)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin
        if name in required:
            cell.fill = required_fill
        elif name in system:
            cell.fill = system_fill
        else:
            cell.fill = optional_fill
        ws.column_dimensions[get_column_letter(col)].width = 18

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["O"].width = 36
    ws.column_dimensions["P"].width = 36
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"

    dv = DataValidation(type="list", formula1='"待处理,成功,失败"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("N2:N200")

    help_ws["A1"] = "Shopify 店铺注册表填写说明"
    help_ws["A1"].font = Font(bold=True, size=14)
    help_ws.merge_cells("A1:C1")
    for r, row in enumerate(HELP_ROWS, 3):
        for c, val in enumerate(row, 1):
            cell = help_ws.cell(r, c, val)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            if r == 3:
                cell.font = header_font
                cell.fill = header_fill
    help_ws.column_dimensions["A"].width = 22
    help_ws.column_dimensions["B"].width = 12
    help_ws.column_dimensions["C"].width = 80
    help_ws.row_dimensions[3].height = 22

    _write_policy_sheet(wb)
    wb.save(excel_path)


def _write_policy_sheet(wb):
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    if SHEET_POLICY in wb.sheetnames:
        ws = wb[SHEET_POLICY]
        ws.delete_rows(1, ws.max_row or 1)
    else:
        ws = wb.create_sheet(SHEET_POLICY)
    ws.cell(1, 1, "类型").font = header_font
    ws.cell(1, 1).fill = header_fill
    ws.cell(1, 2, "文案").font = header_font
    ws.cell(1, 2).fill = header_fill
    for i, (kind, body) in enumerate(SAMPLE_POLICIES, 2):
        ws.cell(i, 1, kind)
        cell = ws.cell(i, 2, body)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 72
    note_row = len(SAMPLE_POLICIES) + 3
    ws.cell(note_row, 1, "说明")
    ws.cell(note_row, 2, (
        "类型必须填：Shipping policy / Legal notice / Terms of sale。"
        "文案可用 {shop_name} {email}。每个店随机抽一条。"
        "若「注册信息」里填了运费政策/法律声明/销售条款，则优先用那一格。"
    ))
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 88


def ensure_policy_sheet(excel_path: str):
    try:
        wb = load_workbook(excel_path)
    except Exception:
        return
    if SHEET_POLICY in wb.sheetnames:
        ws = wb[SHEET_POLICY]
        if (ws.max_row or 0) > 1:
            return
    _write_policy_sheet(wb)
    _save_workbook(wb, excel_path)
    print("已写入「政策模板」工作表，可自行增改文案。")


def pick_policy_text(excel_path: str, kind: str, shop_name: str, email: str, row_override: str = "") -> str:
    """本行有文案用本行，否则从「政策模板」随机抽一条。"""
    text = (row_override or "").strip()
    if not text and excel_path and Path(excel_path).exists():
        try:
            wb = load_workbook(excel_path, data_only=True)
            if SHEET_POLICY in wb.sheetnames:
                options = []
                for row in wb[SHEET_POLICY].iter_rows(min_row=2, max_row=wb[SHEET_POLICY].max_row):
                    k = _cell_str(row[0].value)
                    body = _cell_str(row[1].value if len(row) > 1 else "")
                    if k.lower() == kind.lower() and body and not k.startswith("说明"):
                        options.append(body)
                if options:
                    text = secrets.choice(options)
        except Exception:
            text = ""
    if not text:
        for k, body in SAMPLE_POLICIES:
            if k.lower() == kind.lower():
                text = body
                break
    return (
        (text or "")
        .replace("{shop_name}", shop_name or "")
        .replace("{email}", email or "")
        .strip()
    )


def _cell_str(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _setup_complete(item: dict) -> bool:
    notes = item.get("装修状态", "")
    status = item.get("状态", "")
    if status not in ("成功", "success", "SUCCESS"):
        return False
    if "失败" in notes:
        return False
    return "退货规则已保存" in notes and "书面政策已发布" in notes


def read_registration_data(excel_path: str):
    """读取待处理行。仅跳过「注册+装修」都已成功的行。"""
    ensure_template(excel_path)
    try:
        wb = load_workbook(excel_path, data_only=True)
    except FileNotFoundError:
        print(f"错误：找不到 Excel 文件 '{excel_path}'。")
        return []
    except Exception as e:
        print(f"读取 Excel 失败：{e}")
        return []

    if SHEET_DATA not in wb.sheetnames:
        print(f"错误：Excel 中没有名为「{SHEET_DATA}」的工作表。")
        return []

    ws = wb[SHEET_DATA]
    headers = [_cell_str(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    records = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        if all(c.value is None or str(c.value).strip() == "" for c in row):
            continue
        item = {headers[i]: _cell_str(row[i].value) if i < len(headers) else "" for i in range(len(headers))}
        item["_excel_row"] = row[0].row
        if _setup_complete(item):
            print(f"跳过已完成店铺：{item.get('店铺名') or item.get('邮箱')}")
            continue
        missing = [f for f in REQUIRED_FIELDS if not item.get(f)]
        if missing:
            print(f"第 {item['_excel_row']} 行缺少必填项 {missing}，已跳过。")
            continue
        records.append(item)

    print(f"待处理 {len(records)} 条（含已注册但未完成装修的店铺）。")
    return records


def write_result(excel_path: str, excel_row: int, result: dict):
    """把单条结果写回「注册信息」表。"""
    wb = load_workbook(excel_path)
    ws = wb[SHEET_DATA]
    headers = [_cell_str(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    mapping = {name: idx + 1 for idx, name in enumerate(headers)}

    status = "成功" if result.get("status") == "success" else "失败"
    updates = {
        "状态": status,
        "Shopify密码": result.get("password") or "",
        "店铺后台": result.get("admin_url") or "",
        "装修状态": result.get("setup_notes") or "",
        "失败原因": "" if status == "成功" else (result.get("error") or ""),
        "完成时间": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }
    for key, value in updates.items():
        col = mapping.get(key)
        if not col:
            continue
        cell = ws.cell(excel_row, col)
        if key == "Shopify密码" and cell.value and not value:
            continue
        if key == "Shopify密码" and cell.value and value:
            # 表格里已有密码则保留用户填写的
            if _cell_str(cell.value):
                continue
        cell.value = value or None

    _save_workbook(wb, excel_path)


def _save_workbook(wb, excel_path: str, retries: int = 6):
    last_error = None
    for i in range(1, retries + 1):
        try:
            wb.save(excel_path)
            return
        except PermissionError as e:
            last_error = e
            print(f"无法写入 {excel_path}（可能正被 Excel 打开），2 秒后重试... ({i}/{retries})")
            sleep(2)
    fallback = str(Path(excel_path).with_name(Path(excel_path).stem + "_result.xlsx"))
    wb.save(fallback)
    print(
        f"原表仍被占用，结果已写到 {fallback}。"
        f"请先关闭 Excel 中的 {excel_path} 再运行，避免下次还写不回去。"
        f" 原始错误：{last_error}"
    )
